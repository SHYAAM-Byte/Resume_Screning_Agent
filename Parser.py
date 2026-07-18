"""
parser.py
---------
Resume file parsing utilities.

Scope of this file (and ONLY this file):
    - Read PDF and DOCX resume files
    - Extract clean, plain text from them
    - Fail safely and predictably on corrupted, encrypted, or unsupported files

This file contains NO AI/LLM logic and NO screening logic. It only turns
raw resume files (paths or file-like objects, e.g. Streamlit's
UploadedFile) into plain text for other modules to consume.
"""

from __future__ import annotations

import os
from typing import BinaryIO, Union

import pdfplumber
from PyPDF2 import PdfReader
from PyPDF2.errors import PdfReadError
from docx import Document
from docx.opc.exceptions import PackageNotFoundError


# ---------------------------------------------------------------------------
# Custom exceptions
# ---------------------------------------------------------------------------
class ResumeParsingError(Exception):
    """Base exception for all resume parsing failures."""


class UnsupportedFileTypeError(ResumeParsingError):
    """Raised when a file extension is not supported by the parser."""


class CorruptedFileError(ResumeParsingError):
    """Raised when a file cannot be read because it is corrupted, malformed, or encrypted."""


class EmptyDocumentError(ResumeParsingError):
    """Raised when a file opens successfully but contains no extractable text."""


FileInput = Union[str, "os.PathLike[str]", BinaryIO]


# ---------------------------------------------------------------------------
# Internal helpers
# ---------------------------------------------------------------------------
def _get_file_name(file: FileInput) -> str:
    """Best-effort resolution of a display name for a path or file-like object."""
    name = getattr(file, "name", None)
    if name:
        return os.path.basename(str(name))
    if isinstance(file, (str, os.PathLike)):
        return os.path.basename(str(file))
    return "uploaded_file"


def _get_extension(file: FileInput) -> str:
    """Resolve the lowercase file extension (without the dot)."""
    _, ext = os.path.splitext(_get_file_name(file))
    return ext.lower().lstrip(".")


def _as_readable_stream(file: FileInput) -> BinaryIO:
    """
    Normalize the input into a seekable binary stream.

    Streamlit's UploadedFile (and similar file-like objects) are already
    stream-like, so we just rewind them. Plain paths are opened directly.
    """
    if isinstance(file, (str, os.PathLike)):
        return open(file, "rb")

    if hasattr(file, "seek"):
        try:
            file.seek(0)
        except (OSError, ValueError):
            pass  # Some streams don't support seeking; proceed as-is.
    return file  # type: ignore[return-value]


def _close_if_owned(stream: BinaryIO, original_file: FileInput) -> None:
    """Close the stream only if we opened it ourselves (i.e., original was a path)."""
    if isinstance(original_file, (str, os.PathLike)):
        try:
            stream.close()
        except Exception:
            pass  # Best-effort cleanup; never let this mask the real error.


# ---------------------------------------------------------------------------
# PDF extraction
# ---------------------------------------------------------------------------
def extract_text_from_pdf(file: FileInput) -> str:
    """
    Extract plain text from a PDF file.

    Tries `pdfplumber` first and falls back to `PyPDF2` if that fails,
    since different libraries tolerate different malformed PDF structures.

    Args:
        file: A file path or a file-like object (e.g. Streamlit UploadedFile).

    Returns:
        Extracted plain text.

    Raises:
        CorruptedFileError: If the PDF cannot be opened or is encrypted.
        EmptyDocumentError: If the PDF opens but yields no extractable text
            (e.g. a scanned/image-only PDF with no text layer).
    """
    file_name = _get_file_name(file)
    stream = _as_readable_stream(file)
    text_chunks: list[str] = []

    try:
        try:
            with pdfplumber.open(stream) as pdf:
                for page in pdf.pages:
                    page_text = page.extract_text() or ""
                    if page_text:
                        text_chunks.append(page_text)
        except Exception as primary_error:
            # Fallback attempt with PyPDF2
            try:
                if hasattr(stream, "seek"):
                    stream.seek(0)
                reader = PdfReader(stream)
            except Exception as fallback_open_error:
                raise CorruptedFileError(
                    f"Could not read PDF '{file_name}'. It may be corrupted or "
                    f"not a valid PDF. (pdfplumber error: {primary_error}; "
                    f"PyPDF2 error: {fallback_open_error})"
                ) from fallback_open_error

            if reader.is_encrypted:
                raise CorruptedFileError(
                    f"'{file_name}' is password-protected and cannot be read."
                )

            text_chunks = []
            try:
                for page in reader.pages:
                    page_text = page.extract_text() or ""
                    if page_text:
                        text_chunks.append(page_text)
            except PdfReadError as read_error:
                raise CorruptedFileError(
                    f"Could not read PDF '{file_name}': {read_error}"
                ) from read_error
    finally:
        _close_if_owned(stream, file)

    full_text = "\n".join(text_chunks).strip()
    if not full_text:
        raise EmptyDocumentError(
            f"No extractable text found in '{file_name}'. "
            f"It may be a scanned/image-only PDF without a text layer."
        )
    return full_text


# ---------------------------------------------------------------------------
# DOCX extraction
# ---------------------------------------------------------------------------
def extract_text_from_docx(file: FileInput) -> str:
    """
    Extract plain text from a DOCX file, including paragraphs and table cells.

    Args:
        file: A file path or a file-like object (e.g. Streamlit UploadedFile).

    Returns:
        Extracted plain text.

    Raises:
        CorruptedFileError: If the DOCX cannot be opened or parsed.
        EmptyDocumentError: If the DOCX opens but contains no extractable text.
    """
    file_name = _get_file_name(file)
    stream = _as_readable_stream(file)

    try:
        document = Document(stream)
    except PackageNotFoundError as exc:
        raise CorruptedFileError(
            f"'{file_name}' is not a valid DOCX file or is corrupted."
        ) from exc
    except Exception as exc:
        raise CorruptedFileError(f"Could not read DOCX '{file_name}': {exc}") from exc
    finally:
        _close_if_owned(stream, file)

    text_chunks = [p.text for p in document.paragraphs if p.text and p.text.strip()]

    # Tables frequently hold structured resume data (skills, dates, etc.).
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text and cell.text.strip():
                    text_chunks.append(cell.text.strip())

    full_text = "\n".join(text_chunks).strip()
    if not full_text:
        raise EmptyDocumentError(f"No extractable text found in '{file_name}'.")
    return full_text


# ---------------------------------------------------------------------------
# Public dispatcher
# ---------------------------------------------------------------------------
def extract_text(file: FileInput) -> str:
    """
    Extract plain text from a resume file, auto-detecting PDF vs DOCX by extension.

    Args:
        file: A file path (str/PathLike) or a file-like object with a
            `.name` attribute (e.g. Streamlit's UploadedFile).

    Returns:
        Extracted plain text.

    Raises:
        UnsupportedFileTypeError: If the file extension is not .pdf or .docx.
        CorruptedFileError: If the file exists but cannot be parsed.
        EmptyDocumentError: If the file parses but contains no usable text.
    """
    extension = _get_extension(file)

    if extension == "pdf":
        return extract_text_from_pdf(file)
    if extension == "docx":
        return extract_text_from_docx(file)

    raise UnsupportedFileTypeError(
        f"Unsupported file type '.{extension}'. Only PDF and DOCX are supported."
    )